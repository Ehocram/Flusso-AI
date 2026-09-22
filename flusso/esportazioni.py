"""Esportazione in Excel dell'elenco progetti.

Un foglio unico, una riga per progetto, con le colonne della scheda: anagrafica,
analisi, costi, copertura di budget e compliance. Il file contiene esattamente il
sottoinsieme filtrato a video — la selezione arriva gia' fatta dalla vista, qui
non si filtra nulla — e in testa riporta i filtri applicati, cosi' il foglio resta
leggibile anche mesi dopo, fuori dall'applicazione.
"""

from django.http import HttpResponse
from django.utils import timezone
from django.utils.text import slugify

EURO = '#,##0.00 "€"'
DATA = "DD/MM/YYYY"
DATA_ORA = "DD/MM/YYYY HH:MM"

# (intestazione, come si ricava il valore, formato numerico | None)
COLONNE = [
    ("ID", lambda r: r.codice, None),
    ("Tipo", lambda r: r.tipo_breve, None),
    ("Progetto/Attività", lambda r: r.get_natura_display(), None),
    ("Stato", lambda r: r.stato_label, None),
    ("Titolo", lambda r: r.titolo, None),
    ("Funzione richiedente", lambda r: r.get_funzione_display(), None),
    ("Owner", lambda r: (r.proponente.get_full_name() or r.proponente.username) if r.proponente else "", None),
    ("Priorità owner", lambda r: r.get_priorita_display(), None),
    ("Priorità IT", lambda r: r.get_priorita_it_display() if r.priorita_it else "", None),
    ("Entity", lambda r: r.get_entity_display() if r.entity else "", None),
    ("Tipo soluzione", lambda r: r.tipo_soluzione, None),
    ("Entro quando serve", lambda r: r.data_necessita, DATA),
    ("Effort (ore)", lambda r: r.effort_ore, "0"),
    ("Inizio lavori", lambda r: r.data_inizio, DATA),
    ("Consegna prevista", lambda r: r.data_consegna_prevista, DATA),
    ("SAL %", lambda r: r.sal, "0"),
    ("Costo token annuo (€)", lambda r: r.costo_token_annuo_totale, EURO),
    ("Altri costi (€)", lambda r: r.altri_costi, EURO),
    ("Costo owner (€)", lambda r: r.costo_owner, EURO),
    ("Costo Application (€)", lambda r: r.costo_application, EURO),
    ("Costo IT Operation (€)", lambda r: r.costo_it_operation, EURO),
    ("Costo di progetto (€)", lambda r: r.costo_progetto_stimato, EURO),
    ("Costo iniziativa (€)", lambda r: r.costo_iniziativa, EURO),
    ("Budget IT", lambda r: r.get_budget_it_display() if r.budget_it else "", None),
    ("Copertura", lambda r: r.get_esito_budget_display() if r.esito_budget else "Da definire", None),
    ("Beneficio economico atteso (€)", lambda r: r.saving_economico, EURO),
    ("Incremento qualitativo (%)", lambda r: r.incremento_qualitativo, "0.00"),
    ("Incremento efficienza (%)", lambda r: r.incremento_efficienza, "0.00"),
    ("AI Act", lambda r: _rischio(r, "AIACT"), None),
    ("NIS2", lambda r: _rischio(r, "NIS2"), None),
    ("GDPR", lambda r: _rischio(r, "GDPR"), None),
    ("Validazioni", lambda r: f"{r.rischi_validati_n}/3", None),
    ("Creata il", lambda r: r.creata_il, DATA_ORA),
    ("Ultimo aggiornamento", lambda r: r.aggiornata_il, DATA_ORA),
]


def _rischio(richiesta, tipo):
    """Categoria della dimensione di rischio, con lo stato fra parentesi.

    Legge le classificazioni gia' presenti: l'export non crea righe mancanti
    (a differenza di `lista_rischi`, che le materializza).
    """
    for c in richiesta.classificazioni.all():
        if c.tipo == tipo:
            return f"{c.categoria_label} ({c.get_stato_display()})"
    return "Non valutato"


def _valore(richiesta, estrai):
    valore = estrai(richiesta)
    if valore is None:
        return None
    if hasattr(valore, "utcoffset") and valore.utcoffset() is not None:
        return timezone.localtime(valore).replace(tzinfo=None)  # Excel non regge i fusi
    if hasattr(valore, "quantize"):  # Decimal
        return float(valore)
    return valore


def descrizione_filtri(filtri) -> str:
    """Riga leggibile con i filtri attivi, da mettere in testa al foglio."""
    voci = []
    voci.append("Tipo: " + (filtri.get("tipo_label") or "tutti"))
    if filtri.get("stato_label"):
        voci.append("Stato: " + filtri["stato_label"])
    if filtri.get("funzione_label"):
        voci.append("Funzione: " + filtri["funzione_label"])
    if filtri.get("cerca"):
        voci.append(f"Ricerca: «{filtri['cerca']}»")
    if filtri.get("natura_label"):
        voci.append("Solo: " + filtri["natura_label"])
    if filtri.get("esclude_bozze"):
        voci.append("bozze escluse")
    return " · ".join(voci)


def nome_file(filtri) -> str:
    parti = ["progetti", slugify(filtri.get("tipo_label") or "tutti")]
    if filtri.get("stato_label"):
        parti.append(slugify(filtri["stato_label"]))
    parti.append(timezone.localdate().strftime("%Y%m%d"))
    return "-".join(p for p in parti if p) + ".xlsx"


def risposta_excel(richieste, filtri) -> HttpResponse:
    """Costruisce il workbook e lo restituisce come allegato .xlsx."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Progetti"

    ws["A1"] = "Progetti Digital Transformation — ISEO Group"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = descrizione_filtri(filtri)
    ws["A3"] = (f"{len(richieste)} progetti · esportato il "
                f"{timezone.localtime().strftime('%d/%m/%Y %H:%M')}")
    for cella in ("A2", "A3"):
        ws[cella].font = Font(color="666666")

    riga_intestazione = 5
    testata = Font(bold=True, color="FFFFFF")
    sfondo = PatternFill("solid", fgColor="C8102E")  # rosso ISEO
    for col, (etichetta, _, _) in enumerate(COLONNE, start=1):
        cella = ws.cell(row=riga_intestazione, column=col, value=etichetta)
        cella.font = testata
        cella.fill = sfondo
        cella.alignment = Alignment(vertical="center", wrap_text=True)

    for i, richiesta in enumerate(richieste, start=riga_intestazione + 1):
        for col, (_, estrai, formato) in enumerate(COLONNE, start=1):
            cella = ws.cell(row=i, column=col, value=_valore(richiesta, estrai))
            if formato:
                cella.number_format = formato

    ultima_riga = riga_intestazione + len(richieste)
    ws.freeze_panes = ws.cell(row=riga_intestazione + 1, column=2)
    ws.auto_filter.ref = (f"A{riga_intestazione}:"
                          f"{get_column_letter(len(COLONNE))}{max(ultima_riga, riga_intestazione)}")
    larghezze = {"Titolo": 46, "Tipo soluzione": 30, "Stato": 26, "Owner": 22}
    for col, (etichetta, _, _) in enumerate(COLONNE, start=1):
        ws.column_dimensions[get_column_letter(col)].width = larghezze.get(etichetta,
                                                                          max(12, min(len(etichetta) + 3, 28)))
    ws.row_dimensions[riga_intestazione].height = 30

    risposta = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    risposta["Content-Disposition"] = f'attachment; filename="{nome_file(filtri)}"'
    wb.save(risposta)
    return risposta


# --- Workbook di budget --------------------------------------------------------
# Rileggere i fogli come sono stati importati: stessi nomi, stesse intestazioni,
# stesse righe. L'import dell'Extra Budget prefissa le chiavi con «xb-», quello
# del Budget no: e' cosi' che si riconosce da quale dei due file arriva un foglio.

VIETATI_NEL_NOME = set('[]:*?/\\')


def _famiglia(foglio) -> str:
    return "xb" if foglio.chiave.startswith("xb-") else "bdg"


def fogli_del_workbook(foglio):
    """I fogli del workbook a cui appartiene questo foglio, nell'ordine originale."""
    from .models import FoglioBudget

    famiglia = _famiglia(foglio)
    fogli = [f for f in FoglioBudget.objects.filter(anno=foglio.anno).order_by("ordine", "nome")
             if _famiglia(f) == famiglia]
    return fogli or [foglio]


def _nome_scheda(nome, usati) -> str:
    """Nome di foglio accettato da Excel: max 31 caratteri, senza []:*?/\\, univoco."""
    pulito = "".join(("-" if c in VIETATI_NEL_NOME else c) for c in str(nome)).strip() or "Foglio"
    pulito = pulito[:31]
    base, n = pulito, 2
    while pulito.lower() in usati:
        coda = f"-{n}"
        pulito = base[:31 - len(coda)] + coda
        n += 1
    usati.add(pulito.lower())
    return pulito


def risposta_workbook_budget(foglio) -> HttpResponse:
    """Scarica il workbook di budget nel formato del file di origine."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    fogli = fogli_del_workbook(foglio)
    wb = Workbook()
    wb.remove(wb.active)
    usati = set()
    for f in fogli:
        ws = wb.create_sheet(title=_nome_scheda(f.nome, usati))
        intestazioni = list(f.intestazioni)
        ws.append(intestazioni)
        for cella in ws[1]:
            cella.font = Font(bold=True)
        n_col = len(intestazioni)
        for riga in f.righe.all():
            ws.append(["" if v is None else v for v in riga.celle(n_col)])
        ws.freeze_panes = "A2"
        for col in range(1, n_col + 1):
            lettera = ws.cell(row=1, column=col).column_letter
            ws.column_dimensions[lettera].width = max(
                12, min(len(str(intestazioni[col - 1] or "")) + 3, 40))

    principale = next((f for f in fogli if f.tipo in ("BUDGET", "EXTRA")), foglio)
    etichetta = "Extra_Budget" if _famiglia(foglio) == "xb" else "Budget"
    risposta = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    risposta["Content-Disposition"] = (
        f'attachment; filename="{etichetta}_{principale.anno}.xlsx"')
    wb.save(risposta)
    return risposta
